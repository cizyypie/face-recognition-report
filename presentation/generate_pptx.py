import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen dimensions (13.333 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Executive Color Palette
    c_navy       = RGBColor(15, 23, 42)       # Slate 900 #0f172a
    c_blue       = RGBColor(29, 78, 216)      # Blue 700 #1d4ed8
    c_blue_light = RGBColor(37, 99, 235)      # Blue 600 #2563eb
    c_sky        = RGBColor(2, 132, 199)      # Sky 600 #0284c7
    c_dark       = RGBColor(51, 65, 85)       # Slate 700 #334155
    c_gray       = RGBColor(100, 116, 139)    # Slate 500 #64748b
    c_card_bg    = RGBColor(255, 255, 255)    # Clean White #ffffff
    c_border     = RGBColor(226, 232, 240)    # Slate 200 #e2e8f0
    c_white      = RGBColor(255, 255, 255)
    c_page_bg    = RGBColor(248, 250, 252)    # Slate 50 #f8fafc
    c_header_bar = RGBColor(238, 242, 246)
    
    # Semantic accent colors
    c_green      = RGBColor(21, 128, 61)      # Green 700 #15803d
    c_green_bg   = RGBColor(240, 253, 244)    # Green 50 #f0fdf4
    c_green_bdr  = RGBColor(187, 247, 208)    # Green 200 #bbf7d0
    
    c_red        = RGBColor(185, 28, 28)      # Red 700 #b91c1c
    c_red_bg     = RGBColor(254, 242, 242)    # Red 50 #fef2f2
    c_red_bdr    = RGBColor(254, 202, 202)    # Red 200 #fecaca
    
    c_amber      = RGBColor(180, 83, 9)       # Amber 700 #b45309
    c_amber_bg   = RGBColor(254, 243, 199)    # Amber 100 #fef3c7
    c_amber_bdr  = RGBColor(253, 230, 138)    # Amber 200 #fde68a

    def apply_slide_background(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = c_page_bg
        bg.line.fill.background()

    def add_header(slide, title_text, category_text="WAHANA SWB APP | BRANCH: FACE-ATTENDANCE-HUMAN", slide_num=None):
        apply_slide_background(slide)

        # Category / Breadcrumb
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.42), Inches(10.5), Inches(0.3))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Segoe UI"
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = c_blue

        # Main slide title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(10.5), Inches(0.65))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_right = tf_title.margin_top = tf_title.margin_bottom = 0
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.name = "Segoe UI"
        p_title.font.size = Pt(23)
        p_title.font.bold = True
        p_title.font.color.rgb = c_navy

        # Subtle separator line
        sep = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.02))
        sep.fill.solid()
        sep.fill.fore_color.rgb = c_border
        sep.line.fill.background()

        # Optional slide number footer
        if slide_num:
            f_box = slide.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.0), Inches(0.3))
            tf_f = f_box.text_frame
            p_f = tf_f.paragraphs[0]
            p_f.alignment = PP_ALIGN.RIGHT
            p_f.text = f"{slide_num} / 12"
            p_f.font.name = "Segoe UI"
            p_f.font.size = Pt(10)
            p_f.font.color.rgb = c_gray

    def add_card(slide, left, top, width, height, title, items, badge=""):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_card_bg
        shape.line.color.rgb = c_border
        shape.line.width = Pt(1)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.2)

        start_idx = 0
        if title:
            p_t = tf.paragraphs[0]
            p_t.text = title
            p_t.font.name = "Segoe UI"
            p_t.font.size = Pt(15.5)
            p_t.font.bold = True
            p_t.font.color.rgb = c_blue
            p_t.space_after = Pt(8)
            start_idx = 1

        for i, item in enumerate(items):
            p = tf.add_paragraph() if (title or i > 0) else tf.paragraphs[0]
            p.text = "•  " + item
            p.font.name = "Segoe UI"
            p.font.size = Pt(12)
            p.font.color.rgb = c_dark
            p.space_after = Pt(6)
        return shape

    def add_placeholder_box(slide, left, top, width, height, label_text):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = c_page_bg
        shape.line.color.rgb = c_border
        shape.line.width = Pt(1.5)
        
        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        p = tf.paragraphs[0]
        p.text = label_text
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Segoe UI"
        p.font.size = Pt(12.5)
        p.font.bold = True
        p.font.color.rgb = c_gray
        return shape

    def draw_box(slide, left, top, width, height, title, subtitle="", bg_col=c_card_bg, border_col=c_border, text_col=c_navy, title_size=11, sub_size=9.5):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_col
        shape.line.color.rgb = border_col
        shape.line.width = Pt(1.2)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = Inches(0.12)
        tf.margin_right = Inches(0.12)
        tf.margin_top = Inches(0.08)
        tf.margin_bottom = Inches(0.08)

        p = tf.paragraphs[0]
        p.text = title
        p.alignment = PP_ALIGN.CENTER
        p.font.name = "Segoe UI"
        p.font.size = Pt(title_size)
        p.font.bold = True
        p.font.color.rgb = text_col

        if subtitle:
            p.space_after = Pt(2)
            p2 = tf.add_paragraph()
            p2.text = subtitle
            p2.alignment = PP_ALIGN.CENTER
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(sub_size)
            p2.font.color.rgb = c_gray
        return shape

    def draw_arrow(slide, left, top, width, height, col=c_blue_light):
        shape = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = col
        shape.line.fill.background()
        return shape

    def draw_down_arrow(slide, left, top, width, height, col=c_blue_light):
        shape = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = col
        shape.line.fill.background()
        return shape

    # -------------------------------------------------------------
    # SLIDE 1: TITLE SLIDE
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(241, 245, 249)
    bg1.line.fill.background()

    # Left accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(2.2), Inches(0.15), Inches(3.2))
    bar.fill.solid()
    bar.fill.fore_color.rgb = c_blue
    bar.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.5), Inches(2.1), Inches(10.5), Inches(3.4))
    tf1 = t_box.text_frame
    tf1.word_wrap = True
    
    p0 = tf1.paragraphs[0]
    p0.text = "WAHANA SWB APP  |  MODUL PRESENSI BIOMETRIK"
    p0.font.name = "Segoe UI"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = c_blue
    p0.space_after = Pt(12)

    p1 = tf1.add_paragraph()
    p1.text = "Face Recognition Absensi"
    p1.font.name = "Segoe UI"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = c_navy
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Laporan Penambahan Fitur, Optimasi Performa & Refaktorisasi Arsitektur"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(18)
    p2.font.color.rgb = c_dark
    p2.space_after = Pt(6)

    p3 = tf1.add_paragraph()
    p3.text = "Branch: face-attendance-human  |  Periode: 30 September - 09 Oktober 2026"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(13)
    p3.font.color.rgb = c_gray

    # -------------------------------------------------------------
    # SLIDE 2: LATAR BELAKANG & MASALAH
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Latar Belakang & Masalah Utama", slide_num=2)
    add_card(s2, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Latar Belakang & Kebutuhan Bisnis", [
                 "Presensi nirsentuh berbasis Progressive Web App (PWA) mobile.",
                 "Pencegahan fraud titip absen (buddy punching) pada absensi QR Code.",
                 "Validasi lokasi fisik GPS karyawan secara server-side terintegrasi.",
                 "Kebutuhan otentikasi biometrik yang aman tanpa mengorbankan privasi data."
             ])
    add_card(s2, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Kendala Teknis Sebelum Refaktor", [
                 "Library @vladmandic/human membawa monolitik pipeline ML dalam satu bundle.",
                 "Ukuran bundle puluhan ribu baris menghasilkan cold-start sangat lambat di browser.",
                 "Penurunan frame-rate (lag) parah pada perangkat Android spesifikasi standar.",
                 "Kerentanan terhadap serangan manipulasi foto cetak, layar digital, dan request API."
             ])

    # -------------------------------------------------------------
    # SLIDE 3: PIVOT ARSITEKTUR
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Pivot Arsitektur: Solusi Ringan & Responsif", slide_num=3)
    add_card(s3, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Strategi Pemisahan Komponen (Decoupling)", [
                 "Menanggalkan dependency @vladmandic/human sepenuhnya dari codebase PWA.",
                 "Client PWA beralih ke @vladmandic/face-api khusus inferensi landmark lokal.",
                 "Deteksi wajah berat dialihkan ke sidecar microservice berbasis model YuNet ONNX.",
                 "Ukuran bundle JavaScript turun drastis, rendering instan, dan framerate stabil."
             ])

    # Ilustrasi Skema Pivot Kanan
    pivot_card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1))
    pivot_card.fill.solid()
    pivot_card.fill.fore_color.rgb = c_card_bg
    pivot_card.line.color.rgb = c_border
    pivot_card.line.width = Pt(1)

    p_lbl = s3.shapes.add_textbox(Inches(7.1), Inches(1.85), Inches(5.2), Inches(0.35))
    p_lbl.text_frame.margin_left = p_lbl.text_frame.margin_top = 0
    p_lbl.text_frame.paragraphs[0].text = "SKEMA REFAKTORISASI ENGINE BIOMETRIK"
    p_lbl.text_frame.paragraphs[0].font.name = "Segoe UI"
    p_lbl.text_frame.paragraphs[0].font.size = Pt(12)
    p_lbl.text_frame.paragraphs[0].font.bold = True
    p_lbl.text_frame.paragraphs[0].font.color.rgb = c_blue

    # Sebelum (Merah)
    draw_box(s3, Inches(7.15), Inches(2.35), Inches(5.05), Inches(1.6), 
             "SEBELUM REFAKTOR: MONOLITIK CLIENT", 
             "Library: @vladmandic/human\nBundle > 5 MB | Beban CPU Tinggi | Cold Start Lambat di Mobile", 
             bg_col=c_red_bg, border_col=c_red_bdr, text_col=c_red, title_size=12, sub_size=10.5)

    draw_down_arrow(s3, Inches(9.45), Inches(4.08), Inches(0.45), Inches(0.35), col=c_blue_light)

    # Sesudah (Hijau)
    draw_box(s3, Inches(7.15), Inches(4.55), Inches(2.4), Inches(1.85), 
             "CLIENT (RINGAN)", 
             "@vladmandic/face-api\n+ onnx MiniFASNet\n\nInferensi Landmark Cepat & UI Responsif", 
             bg_col=c_green_bg, border_col=c_green_bdr, text_col=c_green, title_size=11, sub_size=9.5)

    draw_box(s3, Inches(9.8), Inches(4.55), Inches(2.4), Inches(1.85), 
             "SIDECAR (SERVER)", 
             "Node.js Microservice\n+ Model YuNet ONNX\n\nDeteksi Wajah Dominan & Verifikasi Ulang", 
             bg_col=c_green_bg, border_col=c_green_bdr, text_col=c_green, title_size=11, sub_size=9.5)

    # -------------------------------------------------------------
    # SLIDE 4: DIAGRAM ARSITEKTUR 3-TIER
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Arsitektur Sistem Terkini (3-Tier)", slide_num=4)

    # Tier 1 Container
    t1_bg = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(3.6), Inches(5.1))
    t1_bg.fill.solid()
    t1_bg.fill.fore_color.rgb = c_card_bg
    t1_bg.line.color.rgb = c_blue
    t1_bg.line.width = Pt(1.5)
    
    t1_t = s4.shapes.add_textbox(Inches(1.0), Inches(1.85), Inches(3.2), Inches(0.4))
    t1_t.text_frame.margin_left = t1_t.text_frame.margin_top = 0
    t1_t.text_frame.paragraphs[0].text = "TIER 1: CLIENT (PWA BROWSER)"
    t1_t.text_frame.paragraphs[0].font.name = "Segoe UI"
    t1_t.text_frame.paragraphs[0].font.size = Pt(13)
    t1_t.text_frame.paragraphs[0].font.bold = True
    t1_t.text_frame.paragraphs[0].font.color.rgb = c_blue

    draw_box(s4, Inches(1.05), Inches(2.45), Inches(3.1), Inches(1.15), "Vue 2 + Quasar Framework", "UI PWA Mobile & Kontrol Webcam", bg_col=c_page_bg, border_col=c_border)
    draw_box(s4, Inches(1.05), Inches(3.75), Inches(3.1), Inches(1.15), "@vladmandic/face-api", "Ekstraksi 68 Landmark & Pose Liveness", bg_col=c_page_bg, border_col=c_border)
    draw_box(s4, Inches(1.05), Inches(5.05), Inches(3.1), Inches(1.15), "onnxruntime-web (WASM)", "MiniFASNetV2-SE Anti-Spoof Lokal", bg_col=c_page_bg, border_col=c_border)

    # Arrow 1 -> 2
    draw_arrow(s4, Inches(4.55), Inches(4.0), Inches(0.4), Inches(0.35), col=c_blue_light)

    # Tier 2 Container
    t2_bg = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(5.1), Inches(1.65), Inches(3.6), Inches(5.1))
    t2_bg.fill.solid()
    t2_bg.fill.fore_color.rgb = c_card_bg
    t2_bg.line.color.rgb = c_blue
    t2_bg.line.width = Pt(1.5)

    t2_t = s4.shapes.add_textbox(Inches(5.3), Inches(1.85), Inches(3.2), Inches(0.4))
    t2_t.text_frame.margin_left = t2_t.text_frame.margin_top = 0
    t2_t.text_frame.paragraphs[0].text = "TIER 2: BACKEND PERL DCAF"
    t2_t.text_frame.paragraphs[0].font.name = "Segoe UI"
    t2_t.text_frame.paragraphs[0].font.size = Pt(13)
    t2_t.text_frame.paragraphs[0].font.bold = True
    t2_t.text_frame.paragraphs[0].font.color.rgb = c_blue

    draw_box(s4, Inches(5.35), Inches(2.45), Inches(3.1), Inches(1.15), "DCAF Controller: Mobile.pm", "Endpoint API Presensi & Resolusi Profil", bg_col=c_page_bg, border_col=c_border)
    draw_box(s4, Inches(5.35), Inches(3.75), Inches(3.1), Inches(1.15), "Same-Pose Face Matcher", "Euclidean Distance vs Database Internal", bg_col=c_page_bg, border_col=c_border)
    draw_box(s4, Inches(5.35), Inches(5.05), Inches(3.1), Inches(1.15), "Single-Use Token Manager", "Penerbitan & Konsumsi Atomik proofNo", bg_col=c_page_bg, border_col=c_border)

    # Arrow 2 -> 3
    draw_arrow(s4, Inches(8.85), Inches(4.0), Inches(0.4), Inches(0.35), col=c_blue_light)

    # Tier 3 Container
    t3_bg = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.4), Inches(1.65), Inches(3.133), Inches(5.1))
    t3_bg.fill.solid()
    t3_bg.fill.fore_color.rgb = c_card_bg
    t3_bg.line.color.rgb = c_blue
    t3_bg.line.width = Pt(1.5)

    t3_t = s4.shapes.add_textbox(Inches(9.55), Inches(1.85), Inches(2.8), Inches(0.4))
    t3_t.text_frame.margin_left = t3_t.text_frame.margin_top = 0
    t3_t.text_frame.paragraphs[0].text = "TIER 3: AI & DATA"
    t3_t.text_frame.paragraphs[0].font.name = "Segoe UI"
    t3_t.text_frame.paragraphs[0].font.size = Pt(13)
    t3_t.text_frame.paragraphs[0].font.bold = True
    t3_t.text_frame.paragraphs[0].font.color.rgb = c_blue

    draw_box(s4, Inches(9.55), Inches(2.45), Inches(2.8), Inches(1.15), "AI Sidecar (:3100)", "YuNet + MiniFASNet", bg_col=c_page_bg, border_col=c_sky)
    draw_box(s4, Inches(9.55), Inches(3.75), Inches(2.8), Inches(1.15), "Redis In-Memory", "Token attendance:proof (TTL 300s)", bg_col=c_amber_bg, border_col=c_amber_bdr, text_col=c_amber)
    draw_box(s4, Inches(9.55), Inches(5.05), Inches(2.8), Inches(1.15), "MySQL Database", "Tabel EACU (Template Biometrik)", bg_col=c_page_bg, border_col=c_border)

    # -------------------------------------------------------------
    # SLIDE 5: DUAL ANTI-SPOOFING PIPELINE
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Dual Anti-Spoofing Pipeline", slide_num=5)
    add_card(s5, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), 
             "1. Client-Side Anti-Spoof (Pre-Capture)", [
                 "Model MiniFASNetV2-SE quantized via onnxruntime-web (WASM).",
                 "Inferensi real-time langsung di browser sebelum penangkapan frame.",
                 "Mendeteksi anomali tekstur mikro, pantulan cahaya, dan moire pattern layar.",
                 "Visual feedback dinamis pada kamera: Lingkaran Merah (Spoof) vs Hijau (Real).",
                 "Mencegah pengiriman payload foto palsu ke jaringan backend sejak dini."
             ])
    add_card(s5, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1), 
             "2. Server-Side Verification (Post-Capture)", [
                 "Foto hasil auto-capture diverifikasi ulang di sidecar face-service.",
                 "Model YuNet mendeteksi lokasi dan bounding box wajah dominan.",
                 "MiniFASNet server-side mengevaluasi probabilitas keaslian fisik wajah.",
                 "Threshold keamanan ganda memastikan integritas data foto presensi.",
                 "Permintaan presensi otomatis ditolak jika terdeteksi foto cetak / layar."
             ])

    # -------------------------------------------------------------
    # SLIDE 6: UI DETEKSI & ANTI-SPOOFING (ILUSTRASI)
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Tampilan UI Presensi & Deteksi Anti-Spoof", slide_num=6)

    # Kolom Kiri: Real Wajah (Hijau)
    box_real = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1))
    box_real.fill.solid()
    box_real.fill.fore_color.rgb = c_green_bg
    box_real.line.color.rgb = c_green_bdr
    box_real.line.width = Pt(1.5)

    lbl_r = s6.shapes.add_textbox(Inches(1.1), Inches(1.9), Inches(5.1), Inches(0.35))
    lbl_r.text_frame.margin_left = lbl_r.text_frame.margin_top = 0
    lbl_r.text_frame.paragraphs[0].text = "SKENARIO 1: WAJAH ASLI (REAL PERSON)"
    lbl_r.text_frame.paragraphs[0].font.name = "Segoe UI"
    lbl_r.text_frame.paragraphs[0].font.size = Pt(14)
    lbl_r.text_frame.paragraphs[0].font.bold = True
    lbl_r.text_frame.paragraphs[0].font.color.rgb = c_green

    # Ilustrasi Ring Kamera Hijau
    ring_g = s6.shapes.add_shape(MSO_SHAPE.OVAL, Inches(2.65), Inches(2.45), Inches(2.0), Inches(2.2))
    ring_g.fill.solid()
    ring_g.fill.fore_color.rgb = c_white
    ring_g.line.color.rgb = c_green
    ring_g.line.width = Pt(3.5)

    face_g = s6.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3.15), Inches(2.85), Inches(1.0), Inches(1.2))
    face_g.fill.solid()
    face_g.fill.fore_color.rgb = RGBColor(220, 252, 231)
    face_g.line.color.rgb = c_green

    desc_g = s6.shapes.add_textbox(Inches(1.1), Inches(4.9), Inches(5.1), Inches(1.6))
    tf_dg = desc_g.text_frame
    tf_dg.word_wrap = True
    tf_dg.margin_left = tf_dg.margin_top = 0
    p_dg = tf_dg.paragraphs[0]
    p_dg.text = "Hasil Analisis: Skor Keaslian Wajah > 0.85\nPanduan Visual: Lingkaran Hijau (Valid)\nStatus Sistem: Melanjutkan ke Pose Tantangan Liveness"
    p_dg.font.name = "Segoe UI"
    p_dg.font.size = Pt(12)
    p_dg.font.bold = True
    p_dg.font.color.rgb = c_green
    p_dg.alignment = PP_ALIGN.CENTER

    # Kolom Kanan: Spoofing (Merah)
    box_spoof = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1))
    box_spoof.fill.solid()
    box_spoof.fill.fore_color.rgb = c_red_bg
    box_spoof.line.color.rgb = c_red_bdr
    box_spoof.line.width = Pt(1.5)

    lbl_s = s6.shapes.add_textbox(Inches(7.133), Inches(1.9), Inches(5.1), Inches(0.35))
    lbl_s.text_frame.margin_left = lbl_s.text_frame.margin_top = 0
    lbl_s.text_frame.paragraphs[0].text = "SKENARIO 2: SERANGAN SPOOF (FOTO / LAYAR)"
    lbl_s.text_frame.paragraphs[0].font.name = "Segoe UI"
    lbl_s.text_frame.paragraphs[0].font.size = Pt(14)
    lbl_s.text_frame.paragraphs[0].font.bold = True
    lbl_s.text_frame.paragraphs[0].font.color.rgb = c_red

    # Ilustrasi Ring Kamera Merah
    ring_r = s6.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.683), Inches(2.45), Inches(2.0), Inches(2.2))
    ring_r.fill.solid()
    ring_r.fill.fore_color.rgb = c_white
    ring_r.line.color.rgb = c_red
    ring_r.line.width = Pt(3.5)

    phone_r = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.233), Inches(2.8), Inches(0.9), Inches(1.3))
    phone_r.fill.solid()
    phone_r.fill.fore_color.rgb = RGBColor(254, 226, 226)
    phone_r.line.color.rgb = c_red

    desc_r = s6.shapes.add_textbox(Inches(7.133), Inches(4.9), Inches(5.1), Inches(1.6))
    tf_dr = desc_r.text_frame
    tf_dr.word_wrap = True
    tf_dr.margin_left = tf_dr.margin_top = 0
    p_dr = tf_dr.paragraphs[0]
    p_dr.text = "Hasil Analisis: Terdeteksi Pola Moire / Layar Ponsel\nPanduan Visual: Lingkaran Merah (Peringatan)\nStatus Sistem: Auto-Capture Dibatalkan / Dicegah"
    p_dr.font.name = "Segoe UI"
    p_dr.font.size = Pt(12)
    p_dr.font.bold = True
    p_dr.font.color.rgb = c_red
    p_dr.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------
    # SLIDE 7: PENINGKATAN LIVENESS (3 POSE TANTANGAN)
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Peningkatan Liveness: Pose Tantangan Acak", slide_num=7)

    # 3 Kotak Pose Vertikal yang rapi
    # Pose 1: Head Yaw
    p1_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(3.644), Inches(5.1))
    p1_card.fill.solid()
    p1_card.fill.fore_color.rgb = c_card_bg
    p1_card.line.color.rgb = c_border
    p1_card.line.width = Pt(1)

    p1_t = s7.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(3.244), Inches(0.35))
    p1_t.text_frame.margin_left = p1_t.text_frame.margin_top = 0
    p1_t.text_frame.paragraphs[0].text = "1. HEAD YAW (TOLEH)"
    p1_t.text_frame.paragraphs[0].font.name = "Segoe UI"
    p1_t.text_frame.paragraphs[0].font.size = Pt(14)
    p1_t.text_frame.paragraphs[0].font.bold = True
    p1_t.text_frame.paragraphs[0].font.color.rgb = c_blue

    turn_arrow = s7.shapes.add_shape(MSO_SHAPE.CURVED_RIGHT_ARROW, Inches(1.85), Inches(2.45), Inches(1.5), Inches(1.1))
    turn_arrow.fill.solid()
    turn_arrow.fill.fore_color.rgb = c_blue_light
    turn_arrow.line.fill.background()

    p1_desc = s7.shapes.add_textbox(Inches(1.0), Inches(3.8), Inches(3.244), Inches(2.7))
    p1_desc.text_frame.word_wrap = True
    p1_desc.text_frame.margin_left = p1_desc.text_frame.margin_top = 0
    p1_desc.text_frame.paragraphs[0].text = "Gerakan:\nToleh Kiri / Toleh Kanan\n\nLogika Evaluasi:\n• Rasio simetri horizontal ujung hidung terhadap pupil mata.\n• Pergeseran kontur rahang 3D.\n• Ambang batas toleransi dinamis."
    p1_desc.text_frame.paragraphs[0].font.name = "Segoe UI"
    p1_desc.text_frame.paragraphs[0].font.size = Pt(12)
    p1_desc.text_frame.paragraphs[0].font.color.rgb = c_dark

    # Pose 2: Head Tilt (Pitch)
    p2_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.844), Inches(1.65), Inches(3.644), Inches(5.1))
    p2_card.fill.solid()
    p2_card.fill.fore_color.rgb = c_card_bg
    p2_card.line.color.rgb = c_border
    p2_card.line.width = Pt(1)

    p2_t = s7.shapes.add_textbox(Inches(5.044), Inches(1.9), Inches(3.244), Inches(0.35))
    p2_t.text_frame.margin_left = p2_t.text_frame.margin_top = 0
    p2_t.text_frame.paragraphs[0].text = "2. HEAD TILT (PITCH)"
    p2_t.text_frame.paragraphs[0].font.name = "Segoe UI"
    p2_t.text_frame.paragraphs[0].font.size = Pt(14)
    p2_t.text_frame.paragraphs[0].font.bold = True
    p2_t.text_frame.paragraphs[0].font.color.rgb = c_blue

    tilt_up = s7.shapes.add_shape(MSO_SHAPE.UP_ARROW, Inches(5.8), Inches(2.45), Inches(0.55), Inches(0.55))
    tilt_up.fill.solid()
    tilt_up.fill.fore_color.rgb = c_blue_light
    tilt_up.line.fill.background()

    tilt_down = s7.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.9), Inches(3.0), Inches(0.55), Inches(0.55))
    tilt_down.fill.solid()
    tilt_down.fill.fore_color.rgb = c_blue_light
    tilt_down.line.fill.background()

    p2_desc = s7.shapes.add_textbox(Inches(5.044), Inches(3.8), Inches(3.244), Inches(2.7))
    p2_desc.text_frame.word_wrap = True
    p2_desc.text_frame.margin_left = p2_desc.text_frame.margin_top = 0
    p2_desc.text_frame.paragraphs[0].text = "Gerakan:\nTengok Atas / Tengok Bawah\n\nLogika Evaluasi:\n• Rasio vertikal posisi hidung terhadap jarak mata-ke-dagu.\n• Menggantikan pose mulut terbuka (open mouth) yang kurang presisi."
    p2_desc.text_frame.paragraphs[0].font.name = "Segoe UI"
    p2_desc.text_frame.paragraphs[0].font.size = Pt(12)
    p2_desc.text_frame.paragraphs[0].font.color.rgb = c_dark

    # Pose 3: Smile & Multi-Face
    p3_card = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.888), Inches(1.65), Inches(3.644), Inches(5.1))
    p3_card.fill.solid()
    p3_card.fill.fore_color.rgb = c_card_bg
    p3_card.line.color.rgb = c_border
    p3_card.line.width = Pt(1)

    p3_t = s7.shapes.add_textbox(Inches(9.088), Inches(1.9), Inches(3.244), Inches(0.35))
    p3_t.text_frame.margin_left = p3_t.text_frame.margin_top = 0
    p3_t.text_frame.paragraphs[0].text = "3. SMILE & MULTI-FACE"
    p3_t.text_frame.paragraphs[0].font.name = "Segoe UI"
    p3_t.text_frame.paragraphs[0].font.size = Pt(14)
    p3_t.text_frame.paragraphs[0].font.bold = True
    p3_t.text_frame.paragraphs[0].font.color.rgb = c_blue

    smile_shape = s7.shapes.add_shape(MSO_SHAPE.ARC, Inches(10.15), Inches(2.45), Inches(1.1), Inches(0.9))
    smile_shape.line.color.rgb = c_blue_light
    smile_shape.line.width = Pt(3)

    p3_desc = s7.shapes.add_textbox(Inches(9.088), Inches(3.8), Inches(3.244), Inches(2.7))
    p3_desc.text_frame.word_wrap = True
    p3_desc.text_frame.margin_left = p3_desc.text_frame.margin_top = 0
    p3_desc.text_frame.paragraphs[0].text = "Gerakan:\nSenyum Wajar (Smile)\n\nLogika Evaluasi:\n• Rasio pelebaran sudut bibir terhadap jarak antar pupil mata.\n• Multi-Face Rejection: Frame otomatis batal jika terdeteksi > 1 wajah."
    p3_desc.text_frame.paragraphs[0].font.name = "Segoe UI"
    p3_desc.text_frame.paragraphs[0].font.size = Pt(12)
    p3_desc.text_frame.paragraphs[0].font.color.rgb = c_dark

    # -------------------------------------------------------------
    # SLIDE 8: SERVER-SIDE SAME-POSE MATCHING (ILUSTRASI)
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Server-Side Same-Pose Matching", slide_num=8)
    add_card(s8, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Pencocokan Biometrik 1:1 Aman", [
                 "Zero-Leak Biometrics: Data template biometrik tidak pernah dibocorkan ke browser.",
                 "Pencocokan Pose Spesifik: Vektor live pose dicocokkan spesifik ke template pose pendaftaran yang bersesuaian.",
                 "Contoh: Live turn_left diverifikasi langsung ke enroll photoLeft.",
                 "Validasi Euclidean Distance dihitung di backend Perl dengan threshold ketat.",
                 "Menghindari salah deteksi akibat perbedaan sudut perspektif wajah."
             ])

    # Ilustrasi Skema Pencocokan Pose
    match_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1))
    match_card.fill.solid()
    match_card.fill.fore_color.rgb = c_card_bg
    match_card.line.color.rgb = c_border
    match_card.line.width = Pt(1)

    m_lbl = s8.shapes.add_textbox(Inches(7.133), Inches(1.85), Inches(5.1), Inches(0.35))
    m_lbl.text_frame.margin_left = m_lbl.text_frame.margin_top = 0
    m_lbl.text_frame.paragraphs[0].text = "ILUSTRASI SAME-POSE MATCHING (1:1)"
    m_lbl.text_frame.paragraphs[0].font.name = "Segoe UI"
    m_lbl.text_frame.paragraphs[0].font.size = Pt(12)
    m_lbl.text_frame.paragraphs[0].font.bold = True
    m_lbl.text_frame.paragraphs[0].font.color.rgb = c_blue

    # Live Toleh Kiri vs Enroll PhotoLeft
    draw_box(s8, Inches(7.15), Inches(2.4), Inches(2.1), Inches(1.1), "Live Capture", "Pose: Toleh Kiri\n(turn_left)", bg_col=c_page_bg, border_col=c_blue)
    draw_arrow(s8, Inches(9.35), Inches(2.8), Inches(0.45), Inches(0.3), col=c_blue_light)
    draw_box(s8, Inches(9.9), Inches(2.4), Inches(2.2), Inches(1.1), "Enroll Template", "Pose: photoLeft\n(Database EACU)", bg_col=c_green_bg, border_col=c_green_bdr, text_col=c_green)

    # Live Senyum vs Enroll PhotoSmile
    draw_box(s8, Inches(7.15), Inches(3.8), Inches(2.1), Inches(1.1), "Live Capture", "Pose: Senyum\n(smile)", bg_col=c_page_bg, border_col=c_blue)
    draw_arrow(s8, Inches(9.35), Inches(4.2), Inches(0.45), Inches(0.3), col=c_blue_light)
    draw_box(s8, Inches(9.9), Inches(3.8), Inches(2.2), Inches(1.1), "Enroll Template", "Pose: photoSmile\n(Database EACU)", bg_col=c_green_bg, border_col=c_green_bdr, text_col=c_green)

    draw_box(s8, Inches(7.15), Inches(5.25), Inches(4.95), Inches(1.1), 
             "BACKEND PERL EUCLIDEAN DISTANCE EVALUATION", 
             "Jarak Vektor <= Threshold (0.45)  ->  Identitas Terkonfirmasi Sah", 
             bg_col=c_green_bg, border_col=c_green_bdr, text_col=c_green, title_size=11, sub_size=10)

    # -------------------------------------------------------------
    # SLIDE 9: REGISTRASI WAJAH MULTI-POSE
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Registrasi Biometrik (Face Enrollment)", slide_num=9)
    add_card(s9, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Perekaman Template Wajah Karyawan", [
                 "Halaman registrasi biometrik: PTrIAbsenFaceRecognition.vue.",
                 "Perekaman bertahap 6 pose terarah: Center, Left, Right, Up, Down, Smile.",
                 "Ekstraksi vektor embedding 128 dimensi per pose wajah.",
                 "Penyimpanan terenkripsi di tabel database internal (EACU.AACUATR).",
                 "Template biometrik hanya dapat diakses internal oleh backend controller."
             ])
    add_placeholder_box(s9, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1), 
                        "[AREA SCREENSHOT: Antarmuka Registrasi Wajah Multi-Pose (PTrIAbsenFaceRecognition.vue)]")

    # -------------------------------------------------------------
    # SLIDE 10: PENGAMANAN TRANSAKSI (SINGLE-USE TOKEN ILUSTRASI)
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Pengamanan Transaksi: Single-Use Token", slide_num=10)

    token_bg = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.1))
    token_bg.fill.solid()
    token_bg.fill.fore_color.rgb = c_card_bg
    token_bg.line.color.rgb = c_border
    token_bg.line.width = Pt(1)

    t_lbl = s10.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.1), Inches(0.35))
    t_lbl.text_frame.margin_left = t_lbl.text_frame.margin_top = 0
    t_lbl.text_frame.paragraphs[0].text = "SIKLUS HIDUP TOKEN PROOFNO (PENCEGAHAN REPLAY ATTACK & MANIPULASI API)"
    t_lbl.text_frame.paragraphs[0].font.name = "Segoe UI"
    t_lbl.text_frame.paragraphs[0].font.size = Pt(13)
    t_lbl.text_frame.paragraphs[0].font.bold = True
    t_lbl.text_frame.paragraphs[0].font.color.rgb = c_blue

    # Step 1
    draw_box(s10, Inches(1.1), Inches(2.45), Inches(2.5), Inches(2.4), 
             "1. Verifikasi Wajah Lolos", 
             "\nWebcam lolos liveness.\n\nEndpoint API:\n/checkFaceRecongnition\nmemvalidasi pose & identitas.", 
             bg_col=c_page_bg, border_col=c_blue, title_size=12, sub_size=10)

    draw_arrow(s10, Inches(3.7), Inches(3.5), Inches(0.35), Inches(0.3), col=c_blue_light)

    # Step 2
    draw_box(s10, Inches(4.15), Inches(2.45), Inches(2.5), Inches(2.4), 
             "2. Terbitkan proofNo", 
             "\nBackend men-generate token:\nPROOF-<timestamp>-<hex>\n\nDisimpan sementara di Redis\ndengan TTL 300 Detik.", 
             bg_col=c_amber_bg, border_col=c_amber_bdr, text_col=c_amber, title_size=12, sub_size=10)

    draw_arrow(s10, Inches(6.75), Inches(3.5), Inches(0.35), Inches(0.3), col=c_blue_light)

    # Step 3
    draw_box(s10, Inches(7.2), Inches(2.45), Inches(2.5), Inches(2.4), 
             "3. Submit Presensi", 
             "\nPWA mengirimkan payload:\n• Token proofNo\n• Koordinat GPS terverifikasi\n• Foto bukti presensi\nke /attendanceSet.", 
             bg_col=c_page_bg, border_col=c_blue, title_size=12, sub_size=10)

    draw_arrow(s10, Inches(9.8), Inches(3.5), Inches(0.35), Inches(0.3), col=c_blue_light)

    # Step 4
    draw_box(s10, Inches(10.25), Inches(2.45), Inches(1.95), Inches(2.4), 
             "4. Token Dihapus", 
             "\nToken Redis langsung\ndi-DELETE secara atomik.\n\nTidak dapat digunakan ulang (Replay Rejected).", 
             bg_col=c_red_bg, border_col=c_red_bdr, text_col=c_red, title_size=12, sub_size=10)

    t_note = s10.shapes.add_textbox(Inches(1.1), Inches(5.15), Inches(11.1), Inches(1.2))
    tf_tn = t_note.text_frame
    tf_tn.word_wrap = True
    tf_tn.margin_left = tf_tn.margin_top = 0
    p_tn = tf_tn.paragraphs[0]
    p_tn.text = "Hasil Keamanan: Setiap transaksi kehadiran dijamin unik (idempotent). Serangan replay atau request manipulasi otomatis ditolak oleh backend Perl jika token proofNo telah kedaluwarsa atau sudah terhapus."
    p_tn.font.name = "Segoe UI"
    p_tn.font.size = Pt(12)
    p_tn.font.color.rgb = c_dark

    # -------------------------------------------------------------
    # SLIDE 11: DIAGRAM ALUR END-TO-END
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Alur Presensi End-to-End", slide_num=11)

    seq_bg = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.65), Inches(11.733), Inches(5.1))
    seq_bg.fill.solid()
    seq_bg.fill.fore_color.rgb = c_card_bg
    seq_bg.line.color.rgb = c_border
    seq_bg.line.width = Pt(1)

    sq_lbl = s11.shapes.add_textbox(Inches(1.1), Inches(1.85), Inches(11.1), Inches(0.35))
    sq_lbl.text_frame.margin_left = sq_lbl.text_frame.margin_top = 0
    sq_lbl.text_frame.paragraphs[0].text = "DIAGRAM ALUR PROSES PRESENSI LENGKAP (END-TO-END SEQUENCE PIPELINE)"
    sq_lbl.text_frame.paragraphs[0].font.name = "Segoe UI"
    sq_lbl.text_frame.paragraphs[0].font.size = Pt(13)
    sq_lbl.text_frame.paragraphs[0].font.bold = True
    sq_lbl.text_frame.paragraphs[0].font.color.rgb = c_blue

    steps = [
        ("Akses Halaman", "Inisialisasi webcam &\nmuat model WASM"),
        ("Active Liveness", "Tantangan acak\nYaw, Tilt, atau Smile"),
        ("Dual Anti-Spoof", "Analisis tekstur lokal &\nvalidasi sidecar"),
        ("Check Face API", "Backend Same-Pose\nMatching 1:1"),
        ("Terbit Token", "Simpan proofNo\ndi Redis (TTL 300s)"),
        ("Submit Presensi", "Catat absensi & hapus\ntoken Redis")
    ]

    for idx, (st_title, st_desc) in enumerate(steps):
        x = Inches(1.1 + idx * 1.88)
        draw_box(s11, x, Inches(2.45), Inches(1.68), Inches(2.4), 
                 f"{idx+1}. {st_title}", "\n" + st_desc, bg_col=c_page_bg, border_col=c_blue, title_size=11, sub_size=9.5)
        if idx < len(steps) - 1:
            draw_arrow(s11, x + Inches(1.71), Inches(3.5), Inches(0.14), Inches(0.2), col=c_blue_light)

    res_box = s11.shapes.add_textbox(Inches(1.1), Inches(5.15), Inches(11.1), Inches(1.2))
    tf_rb = res_box.text_frame
    tf_rb.word_wrap = True
    tf_rb.margin_left = tf_rb.margin_top = 0
    p_rb = tf_rb.paragraphs[0]
    p_rb.text = "Hasil: Presensi tercatat sah dengan bukti foto ber-watermark waktu, koordinat GPS terverifikasi server, dan biometrik wajah yang tervalidasi tanpa kebocoran data biometrik ke sisi browser."
    p_rb.font.name = "Segoe UI"
    p_rb.font.size = Pt(12)
    p_rb.font.bold = True
    p_rb.font.color.rgb = c_navy

    # -------------------------------------------------------------
    # SLIDE 12: RINGKASAN & LANGKAH SELANJUTNYA
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Ringkasan Hasil & Rencana Selanjutnya", slide_num=12)
    add_card(s12, Inches(0.8), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Pencapaian Utama (12 Commit)", [
                 "Aplikasi PWA ringan dan responsif tanpa lag pada berbagai smartphone.",
                 "Perlindungan menyeluruh terhadap spoofing layar, foto cetak, dan video rekaman.",
                 "Keamanan biometrik terjamin dengan Zero-Leak dan Same-Pose Matching.",
                 "Pencatatan presensi terlindungi dari manipulasi request API berkat Single-Use Token.",
                 "Pipeline arsitektur modular yang siap di-scale ke lingkungan produksi."
             ])
    add_card(s12, Inches(6.833), Inches(1.65), Inches(5.7), Inches(5.1), 
             "Rencana Selanjutnya (Roadmap)", [
                 "Peningkatan toleransi deteksi pada kondisi pencahayaan rendah (low-light).",
                 "Pengujian komprehensif pada variasi perangkat mobile entry-level.",
                 "Mekanisme caching model lokal yang lebih optimal pada service worker PWA.",
                 "Monitoring metrik latensi inferensi sidecar container di lingkungan production."
             ])

    prs.save("/home/joan/codes/new/swb/doc/face-recognition-presentation.pptx")
    print("SUCCESS: Beautifully polished presentation saved to /home/joan/codes/new/swb/doc/face-recognition-presentation.pptx")

if __name__ == "__main__":
    create_presentation()
